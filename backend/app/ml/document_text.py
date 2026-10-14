"""Text extraction from uploaded resume files (PDF and DOCX).

Files are treated as untrusted input:

* the declared extension must agree with the file's magic bytes;
* DOCX archives are checked for the expected structure and for zip bombs;
* encrypted, empty, malformed or image-only (scanned) PDFs are rejected with
  a clear, user-facing explanation rather than a stack trace.

Text is never logged.
"""

from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field

import fitz  # PyMuPDF
from docx import Document

PDF_MAGIC = b"%PDF-"
ZIP_MAGIC = b"PK\x03\x04"
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024
MIN_TEXT_CHARS = 80


class DocumentError(Exception):
    """A user-facing problem with an uploaded document."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class ExtractedDocument:
    text: str
    content_type: str  # pdf | docx
    page_count: int | None
    method: str
    warnings: list[str] = field(default_factory=list)


_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def normalize_text(text: str, max_chars: int) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = _CONTROL.sub(" ", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t ]+", " ", line).strip() for line in text.split("\n")]
    # Collapse runs of blank lines.
    out: list[str] = []
    for line in lines:
        if line or (out and out[-1]):
            out.append(line)
    return "\n".join(out).strip()[:max_chars]


def detect_type(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in {"pdf", "docx"}:
        raise DocumentError("unsupported_type", "Only PDF and DOCX files are supported.")
    if ext == "pdf" and not data.startswith(PDF_MAGIC):
        raise DocumentError("type_mismatch", "The file does not appear to be a valid PDF.")
    if ext == "docx" and not data.startswith(ZIP_MAGIC):
        raise DocumentError("type_mismatch", "The file does not appear to be a valid DOCX document.")
    return ext


def _extract_pdf(data: bytes) -> tuple[str, int, list[str]]:
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:  # malformed PDF
        raise DocumentError("malformed", "The PDF could not be read. It may be damaged.") from exc
    try:
        if doc.needs_pass or doc.is_encrypted:
            raise DocumentError("encrypted", "The PDF is password-protected. Please upload an unprotected copy.")
        warnings: list[str] = []
        pages = []
        for page in doc:
            pages.append(page.get_text("text", sort=True))
        text = "\n".join(pages)
        image_pages = sum(1 for page in doc if page.get_images() and len(page.get_text("text").strip()) < 20)
        if image_pages:
            warnings.append(
                f"{image_pages} page(s) appear to contain only images; text on those pages could not be read."
            )
        return text, doc.page_count, warnings
    finally:
        doc.close()


def _extract_docx(data: bytes) -> tuple[str, list[str]]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
            if "word/document.xml" not in names:
                raise DocumentError("malformed", "The DOCX file is missing its document body.")
            total = sum(info.file_size for info in archive.infolist())
            if total > MAX_DOCX_UNCOMPRESSED:
                raise DocumentError("too_large", "The DOCX file expands to an unsafe size.")
    except zipfile.BadZipFile as exc:
        raise DocumentError("malformed", "The DOCX file could not be read. It may be damaged.") from exc
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError("malformed", "The DOCX file could not be read. It may be damaged.") from exc
    parts: list[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                value = cell.text.strip()
                if value and value not in cells:
                    cells.append(value)
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts), []


def extract_document(filename: str, data: bytes, max_chars: int = 60_000) -> ExtractedDocument:
    if not data:
        raise DocumentError("empty_file", "The uploaded file is empty.")
    kind = detect_type(filename, data)
    if kind == "pdf":
        raw, page_count, warnings = _extract_pdf(data)
        method = "pymupdf-text"
    else:
        raw, warnings = _extract_docx(data)
        page_count = None
        method = "python-docx"
    text = normalize_text(raw, max_chars)
    if len(text) < MIN_TEXT_CHARS:
        if kind == "pdf":
            raise DocumentError(
                "no_text",
                "No readable text was found. The PDF may be a scanned image. OCR is not enabled in this "
                "deployment — please upload a text-based PDF or DOCX, or enter your skills manually.",
            )
        raise DocumentError("no_text", "The document contains too little text to analyse.")
    if len(raw) > max_chars:
        warnings.append("The document was very long; only the first part was analysed.")
    return ExtractedDocument(text=text, content_type=kind, page_count=page_count, method=method, warnings=warnings)
